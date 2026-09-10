#!/usr/bin/env python3
"""
Fetch 10-year and 30-year government bond yields for six countries from
official sources and merge them into data/yields.json.

Every source is an official publisher (treasury / central bank). All values are
end-of-day. Nothing here is intraday. No API keys required.

Sources
  US  US Treasury daily yield curve CSV          (10 Yr / 30 Yr par yields)
  DE  Deutsche Bundesbank time series API        (Svensson spot, 10.0y / 30.0y)
  UK  Bank of England GLC nominal daily data     (spot curve, 10y / 30y)
  FR  Eurostat Maastricht convergence rate       (10Y only, MONTHLY -- see README)
  CA  Bank of Canada Valet API                   (10y benchmark / long benchmark)
  JP  Japan MoF JGB interest rate CSV            (10Y / 30Y)

Each fetcher returns {"10Y": {iso_date: yield}, "30Y": {...}}.

Run:  python scripts/fetch_yields.py
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import pathlib
import sys
import zipfile

import requests

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA_FILE = ROOT / "data" / "yields.json"

HISTORY_YEARS = 3
MATURITIES = ["10Y", "30Y"]

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)
TIMEOUT = 90

COUNTRIES = {
    "US": {"name": "United States", "bond": "Treasury", "source": "US Treasury"},
    "DE": {"name": "Germany", "bond": "Bund", "source": "Deutsche Bundesbank"},
    "UK": {"name": "United Kingdom", "bond": "Gilt", "source": "Bank of England"},
    "FR": {
        "name": "France",
        "bond": "OAT",
        "source": "Eurostat",
        "freq": "monthly",
        # France has no 30Y source at all; the 10Y stands in, flagged in the UI.
        "only": "10Y",
        "caveat": "No official daily French yield since July 2024. Eurostat 10Y monthly average.",
    },
    "CA": {"name": "Canada", "bond": "GoC benchmark", "source": "Bank of Canada"},
    "JP": {"name": "Japan", "bond": "JGB", "source": "Japan Ministry of Finance"},
}

# Germany is the spread anchor for the whole page.
ANCHOR = "DE"


def get(url: str, **kw) -> requests.Response:
    r = requests.get(url, headers={"User-Agent": UA}, timeout=TIMEOUT, **kw)
    r.raise_for_status()
    return r


def clean(pairs) -> dict[str, float]:
    out = {}
    for date, value in pairs:
        if value in (None, "", ".", "-", "ND", "NaN"):
            continue
        try:
            out[date] = round(float(value), 4)
        except (TypeError, ValueError):
            continue
    return out


# --------------------------------------------------------------------------
# Per-country fetchers
# --------------------------------------------------------------------------

def fetch_us() -> dict[str, dict[str, float]]:
    years = range(dt.date.today().year - HISTORY_YEARS + 1, dt.date.today().year + 1)
    out = {m: [] for m in MATURITIES}
    for year in years:
        url = (
            "https://home.treasury.gov/resource-center/data-chart-center/"
            f"interest-rates/daily-treasury-rates.csv/{year}/all"
            f"?type=daily_treasury_yield_curve&field_tdr_date_value={year}"
            "&page&_format=csv"
        )
        for row in csv.DictReader(io.StringIO(get(url).text)):
            raw = row.get("Date")
            if not raw:
                continue
            date = dt.datetime.strptime(raw.strip(), "%m/%d/%Y").date().isoformat()
            out["10Y"].append((date, row.get("10 Yr")))
            out["30Y"].append((date, row.get("30 Yr")))
    return {m: clean(v) for m, v in out.items()}


def fetch_de() -> dict[str, dict[str, float]]:
    keys = {"10Y": "R10XX", "30Y": "R30XX"}
    out = {}
    for maturity, code in keys.items():
        series = f"D.I.ZST.ZI.EUR.S1311.B.A604.{code}.R.A.A._Z._Z.A"
        url = (
            f"https://api.statistiken.bundesbank.de/rest/download/BBSIS/{series}"
            "?format=csv&lang=en"
        )
        pairs = []
        for row in csv.reader(io.StringIO(get(url).text)):
            if len(row) < 2:
                continue
            date = row[0].strip().strip('"').lstrip("\ufeff")
            if len(date) == 10 and date[4] == "-" and date[7] == "-":
                pairs.append((date, row[1].strip()))
        out[maturity] = clean(pairs)
    return out


def fetch_uk(current_month_only: bool = True) -> dict[str, dict[str, float]]:
    """Bank of England nominal spot curve, delivered as xlsx inside a zip.

    The 'latest' zip holds the current month only, so this merges into an
    accumulating history. Use scripts/backfill_uk.py once to load the archive.
    """
    import openpyxl

    url = (
        "https://www.bankofengland.co.uk/-/media/boe/files/statistics/"
        "yield-curves/latest-yield-curve-data.zip"
        if current_month_only
        else "https://www.bankofengland.co.uk/-/media/boe/files/statistics/"
        "yield-curves/glcnominalddata.zip"
    )
    archive = zipfile.ZipFile(io.BytesIO(get(url).content))
    names = [n for n in archive.namelist() if "Nominal" in n and n.endswith(".xlsx")]
    out = {m: [] for m in MATURITIES}

    for name in names:
        book = openpyxl.load_workbook(
            io.BytesIO(archive.read(name)), read_only=True, data_only=True
        )
        if "4. spot curve" not in book.sheetnames:
            continue
        sheet = book["4. spot curve"]

        columns = {}
        for row in sheet.iter_rows(values_only=True):
            if not row:
                continue
            head = row[0]
            if isinstance(head, str) and head.strip().lower() == "years:":
                for idx, cell in enumerate(row):
                    if not isinstance(cell, (int, float)):
                        continue
                    if abs(cell - 10) < 1e-9:
                        columns["10Y"] = idx
                    elif abs(cell - 30) < 1e-9:
                        columns["30Y"] = idx
                break
        if not columns:
            continue

        for row in sheet.iter_rows(values_only=True):
            if not row or not isinstance(row[0], dt.datetime):
                continue
            date = row[0].date().isoformat()
            for maturity, idx in columns.items():
                if len(row) > idx:
                    out[maturity].append((date, row[idx]))
        book.close()

    return {m: clean(v) for m, v in out.items()}


def fetch_fr() -> dict[str, dict[str, float]]:
    """France, via Eurostat's EMU convergence criterion series (monthly).

    Banque de France stopped publishing OAT rates on 10 July 2024 and no
    official publisher replaced them -- French govvies trade OTC, so the only
    daily signal comes from commercial dealer-quote feeds. Eurostat's Maastricht
    long-term rate is the closest free official substitute: 10-year only, and a
    monthly average rather than a daily close. There is no 30Y equivalent.
    """
    url = (
        "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/"
        "irt_lt_mcby_m"
    )
    since = f"{dt.date.today().year - HISTORY_YEARS}-01"
    payload = get(
        url, params={"format": "JSON", "geo": "FR", "sinceTimePeriod": since}
    ).json()

    periods = payload["dimension"]["time"]["category"]["index"]
    by_position = {position: period for period, position in periods.items()}

    pairs = []
    for flat_index, value in payload.get("value", {}).items():
        period = by_position.get(int(flat_index))
        if period:
            # A monthly average belongs mid-month, not on the 1st.
            pairs.append((f"{period}-15", value))
    return {"10Y": clean(pairs), "30Y": {}}


def fetch_ca() -> dict[str, dict[str, float]]:
    start = (dt.date.today() - dt.timedelta(days=365 * HISTORY_YEARS)).isoformat()
    # Canada has no formal 30Y benchmark; the long-term benchmark is the ~30Y bond.
    keys = {"10Y": "BD.CDN.10YR.DQ.YLD", "30Y": "BD.CDN.LONG.DQ.YLD"}
    out = {}
    for maturity, series in keys.items():
        url = (
            "https://www.bankofcanada.ca/valet/observations/"
            f"{series}/json?start_date={start}"
        )
        payload = get(url).json()
        out[maturity] = clean(
            (obs["d"], obs.get(series, {}).get("v"))
            for obs in payload.get("observations", [])
        )
    return out


def fetch_jp() -> dict[str, dict[str, float]]:
    base = "https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/"
    out = {m: [] for m in MATURITIES}
    for suffix in ("historical/jgbcme_all.csv", "jgbcme.csv"):
        text = get(base + suffix).content.decode("shift_jis", errors="replace")
        columns = {}
        for row in csv.reader(io.StringIO(text)):
            if not row:
                continue
            if row[0].strip() == "Date":
                columns = {m: row.index(m) for m in MATURITIES if m in row}
                continue
            if not columns:
                continue
            try:
                date = dt.datetime.strptime(row[0].strip(), "%Y/%m/%d").date().isoformat()
            except ValueError:
                continue
            for maturity, idx in columns.items():
                if len(row) > idx:
                    out[maturity].append((date, row[idx].strip()))
    return {m: clean(v) for m, v in out.items()}


FETCHERS = {
    "US": fetch_us, "DE": fetch_de, "UK": fetch_uk,
    "FR": fetch_fr, "CA": fetch_ca, "JP": fetch_jp,
}


def main() -> int:
    existing = {}
    if DATA_FILE.exists():
        existing = json.loads(DATA_FILE.read_text()).get("countries", {})

    cutoff = (dt.date.today() - dt.timedelta(days=365 * HISTORY_YEARS + 30)).isoformat()
    countries, failures = {}, []

    for code, meta in COUNTRIES.items():
        prior = existing.get(code, {}).get("series", {})
        # Tolerate the old single-maturity file shape on first upgrade.
        if isinstance(prior, list):
            prior = {"30Y": prior}
        history = {m: {d: v for d, v in prior.get(m, [])} for m in MATURITIES}

        try:
            fresh = FETCHERS[code]()
            for maturity in MATURITIES:
                history[maturity].update(fresh.get(maturity, {}))
            counts = ", ".join(f"{m} +{len(fresh.get(m, {}))}" for m in MATURITIES)
            print(f"  {code}: {counts}")
        except Exception as exc:
            failures.append(code)
            print(f"  {code}: FAILED ({exc.__class__.__name__}: {exc})")

        series, last, last_date = {}, {}, {}
        for maturity in MATURITIES:
            points = sorted((d, v) for d, v in history[maturity].items() if d >= cutoff)
            if not points:
                continue
            series[maturity] = [list(p) for p in points]
            last[maturity] = points[-1][1]
            last_date[maturity] = points[-1][0]

        if not series:
            continue
        countries[code] = {**meta, "series": series, "last": last, "lastDate": last_date}
        print(f"      -> {', '.join(f'{m}={last[m]}%' for m in sorted(last))}")

    if not countries:
        print("No data at all; refusing to overwrite.")
        return 1

    DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
    DATA_FILE.write_text(
        json.dumps(
            {
                "updated": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
                "maturities": MATURITIES,
                "anchor": ANCHOR,
                "note": "End-of-day official data. Not intraday, not investment advice.",
                "stale": failures,
                "countries": countries,
            },
            separators=(",", ":"),
        )
    )
    print(f"\nWrote {DATA_FILE} ({DATA_FILE.stat().st_size / 1024:.0f} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
